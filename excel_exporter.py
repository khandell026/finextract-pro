import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

def export_to_excel(data_rows, output_path):
    """
    Creates a professionally formatted Excel spreadsheet from the data rows.
    Applies custom teal header styles, borders, number formatting, and auto-adjusts widths.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Invoice Register"
    
    # Enable grid lines
    ws.views.sheetView[0].showGridLines = True
    
    # Custom Teal Palette & Typography
    TEAL_HEADER_FILL = PatternFill(start_color="008B9B", end_color="008B9B", fill_type="solid")
    WHITE_BOLD_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    NORMAL_FONT = Font(name="Calibri", size=11, color="333333")
    TOTALS_FONT = Font(name="Calibri", size=11, bold=True, color="000000")
    
    # Alignments
    LEFT_ALIGN = Alignment(horizontal="left", vertical="center")
    RIGHT_ALIGN = Alignment(horizontal="right", vertical="center")
    CENTER_ALIGN = Alignment(horizontal="center", vertical="center")
    
    # Borders
    THIN_BORDER_SIDE = Side(border_style="thin", color="D3D3D3")
    DOUBLE_BORDER_SIDE = Side(border_style="double", color="000000")
    THICK_BORDER_SIDE = Side(border_style="medium", color="000000")
    
    CELL_BORDER = Border(left=THIN_BORDER_SIDE, right=THIN_BORDER_SIDE, top=THIN_BORDER_SIDE, bottom=THIN_BORDER_SIDE)
    HEADER_BORDER = Border(left=THIN_BORDER_SIDE, right=THIN_BORDER_SIDE, top=THICK_BORDER_SIDE, bottom=THICK_BORDER_SIDE)
    TOTALS_BORDER = Border(top=THIN_BORDER_SIDE, bottom=DOUBLE_BORDER_SIDE)

    # Column Headers (exactly matching user's layout starting from column B)
    headers = [
        "File name", "Vendor Name", "Vendor GSTIN", "Buyer Name", "Buyer GSTIN", 
        "Nature of Service", "Invoice Type", "Invoice No.", "Invoice Date", 
        "Taxable Amount", "IGST rate", "CGST rate", "SGST rate", 
        "IGST Amount", "CGST Amount", "SGST Amount", "Total invoice amount"
    ]
    
    # Set header row height
    ws.row_dimensions[2].height = 28
    
    # Write Headers starting at B2
    for col_idx, header in enumerate(headers, start=2): # Col 2 is 'B'
        cell = ws.cell(row=2, column=col_idx)
        cell.value = header
        cell.fill = TEAL_HEADER_FILL
        cell.font = WHITE_BOLD_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = HEADER_BORDER

    # Write Data Rows starting at B3
    current_row = 3
    for row_data in data_rows:
        ws.row_dimensions[current_row].height = 20
        
        # Mapping values to columns B to R
        values = [
            row_data.get("File name", ""),
            row_data.get("Vendor Name", ""),
            row_data.get("Vendor GSTIN", ""),
            row_data.get("Buyer Name", ""),
            row_data.get("Buyer GSTIN", ""),
            row_data.get("Nature of Service", ""),
            row_data.get("Invoice Type", ""),
            row_data.get("Invoice No.", ""),
            row_data.get("Invoice Date", ""),
            row_data.get("Taxable Amount", 0.0),
            row_data.get("IGST rate", 0.0),
            row_data.get("CGST rate", 0.0),
            row_data.get("SGST rate", 0.0),
            row_data.get("IGST Amount", 0.0),
            row_data.get("CGST Amount", 0.0),
            row_data.get("SGST Amount", 0.0),
            row_data.get("Total invoice amount", 0.0)
        ]
        
        for col_idx, val in enumerate(values, start=2):
            cell = ws.cell(row=current_row, column=col_idx)
            cell.value = val
            cell.font = NORMAL_FONT
            cell.border = CELL_BORDER
            
            # Format and align by data type
            if col_idx in [2, 3, 4, 5, 6, 7, 8]: # Text columns (File Name, Vendor, GSTINs, Buyer, Nature of Service, Invoice Type)
                cell.alignment = LEFT_ALIGN
            elif col_idx in [9]: # Invoice No
                cell.alignment = CENTER_ALIGN
            elif col_idx in [10]: # Invoice Date
                cell.alignment = CENTER_ALIGN
            elif col_idx in [11]: # Taxable Amount (formatted as Currency/Decimals)
                cell.number_format = '#,##0.00'
                cell.alignment = RIGHT_ALIGN
            elif col_idx in [12, 13, 14]: # Rates (formatted as clean numbers/percentages)
                cell.number_format = '0'
                cell.alignment = CENTER_ALIGN
            elif col_idx in [15, 16, 17, 18]: # GST Amounts and Grand Total
                cell.number_format = '#,##0.00'
                cell.alignment = RIGHT_ALIGN
                
        current_row += 1

    # Add Summary/Total Row at the end
    if len(data_rows) > 0:
        ws.row_dimensions[current_row].height = 22
        
        # "Total" label in Column G (Nature of Service) or B
        total_cell = ws.cell(row=current_row, column=2)
        total_cell.value = "TOTAL"
        total_cell.font = TOTALS_FONT
        total_cell.border = TOTALS_BORDER
        total_cell.alignment = LEFT_ALIGN
        
        # Apply border to columns B through R on totals row
        for col_idx in range(2, 19):
            cell = ws.cell(row=current_row, column=col_idx)
            cell.border = TOTALS_BORDER
            if col_idx == 2:
                continue
                
            # Excel SUM formulas for numeric columns
            col_letter = get_column_letter(col_idx)
            start_cell = f"{col_letter}3"
            end_cell = f"{col_letter}{current_row-1}"
            
            if col_idx == 11: # Taxable Amount
                cell.value = f"=SUM({start_cell}:{end_cell})"
                cell.number_format = '#,##0.00'
                cell.font = TOTALS_FONT
                cell.alignment = RIGHT_ALIGN
            elif col_idx in [15, 16, 17, 18]: # GST and Grand Total
                cell.value = f"=SUM({start_cell}:{end_cell})"
                cell.number_format = '#,##0.00'
                cell.font = TOTALS_FONT
                cell.alignment = RIGHT_ALIGN
                
    # Auto-adjust column widths with safety margin
    for col in ws.columns:
        col_idx = col[0].column
        if col_idx < 2 or col_idx > 18:
            continue
            
        max_len = 0
        col_letter = get_column_letter(col_idx)
        
        for cell in col:
            val_to_check = str(cell.value or '')
            if cell.number_format == '#,##0.00' and isinstance(cell.value, (int, float)):
                val_to_check = f"{cell.value:,.2f}"
            max_len = max(max_len, len(val_to_check))
            
        # Set generous width
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # Save workbook
    wb.save(output_path)
